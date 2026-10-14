# -*- coding: utf-8 -*-
"""Build editable PowerPoint versions of the ENG 3861 decks from decks.yml.

Same source as build_decks.py (the reveal.js web decks). These .pptx files
are what goes to Canvas: the course is handed off to an instructor who
edits slides in PowerPoint, not in this repo. Rebuild after editing decks.yml:

    ../../.venv/Scripts/python build_pptx.py      (from undergrad/slides/)

To look at every slide as PowerPoint renders it (Windows + PowerPoint):

    powershell -File export_pngs.ps1 -OutDir <some folder>
"""
from __future__ import annotations

import io
import os
import re
import sys
import zipfile
from datetime import datetime
from functools import lru_cache
from pathlib import Path

import frontmatter
import yaml
from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.text.text import _Run
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
COURSE = HERE.parent
REPO = COURSE.parent

TITLE_FONT, BODY_FONT, MONO_FONT = "Arial Black", "Verdana", "Courier New"
C = {k: RGBColor.from_string(v) for k, v in {
    "paper": "FFF7E3", "panel": "FFFFFF", "ink": "1B1035", "soft": "4A3F66",
    "pink": "FF3D7F", "blue": "2E5BFF", "yellow": "FFD23F", "mint": "00B894",
    "orange": "FF7A1A", "termbg": "1B1035", "termfg": "7DFFB3",
    "bar1": "FFD23F", "bar2": "8FE9D3", "bar3": "FFB3CF", "bar0": "B9C8FF"}.items()}
UNIT = {"words": "blue", "images": "pink", "code": "mint", "coda": "orange"}
UNIT_LABELS = {"words": "Words", "images": "Images", "code": "Code", "coda": "Futures"}
DARK_TEXT_UNITS = {"code", "coda"}   # white text is unreadable on mint and orange
W, H = Inches(13.333), Inches(7.5)
MAX_PX = 1920
FIXED_DATE = datetime(2027, 1, 1)   # core properties, so rebuilds don't churn

_TOKEN = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)|\*\*(.+?)\*\*|\*(.+?)\*")


def runs(text: str) -> list[tuple[str, bool, bool, str | None]]:
    """Split inline markup into (text, bold, italic, url) runs. Spaces are kept."""
    text, out, pos = str(text), [], 0
    for m in _TOKEN.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False, False, None))
        if m.group(1) is not None:
            for t, b, i, _ in runs(m.group(1)):   # *italic* inside link text
                out.append((t, b, i, m.group(2)))
        elif m.group(3) is not None:
            out.append((m.group(3), True, False, None))
        else:
            out.append((m.group(4), False, True, None))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False, False, None))
    return out


def plain(text: str) -> str:
    return "".join(r[0] for r in runs(str(text).strip()))


def write(p, text, size, font=BODY_FONT, color="ink", bold=False):
    """Append marked-up text to paragraph p; a "\\n" becomes a line break."""
    for k, line in enumerate(str(text).split("\n")):
        if k:
            p.add_line_break()
        for t, b, i, url in runs(line):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name, f.size, f.bold, f.italic = font, Pt(size), bold or b, i
            f.color.rgb = C[color]
            if url:
                r.hyperlink.address = url


# ---------- measuring text the way PowerPoint will lay it out ----------
# python-pptx's TEXT_TO_FIT_SHAPE only writes <a:normAutofit/>, and PowerPoint
# doesn't shrink anything until someone edits the text. So every text box is
# filled at the largest size from a ladder that measurably fits, using the
# real font files. Autofit stays on so the GTA's later edits still shrink.
FONT_DIR = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
FONT_FILES = {(TITLE_FONT, False): "ariblk.ttf", (TITLE_FONT, True): "ariblk.ttf",
              (BODY_FONT, False): "verdana.ttf", (BODY_FONT, True): "verdanab.ttf",
              (MONO_FONT, False): "cour.ttf", (MONO_FONT, True): "courbd.ttf"}
FALLBACK_EM = {TITLE_FONT: (0.70, 1.42), BODY_FONT: (0.62, 1.22), MONO_FONT: (0.60, 1.15)}
SAFETY = 0.95   # measured text may use at most 95% of a box, each way


@lru_cache(maxsize=None)
def _font(name: str, bold: bool):
    try:
        return ImageFont.truetype(str(FONT_DIR / FONT_FILES[(name, bold)]), 100)
    except (OSError, KeyError):
        return None


def _width(text: str, name: str, bold: bool, size: float) -> float:
    """Advance width in points (estimated if the font file isn't installed)."""
    f = _font(name, bool(bold))
    if f is None:
        return len(text) * FALLBACK_EM.get(name, (0.65, 1.3))[0] * size
    return f.getlength(text) * size / 100


def _line_em(name: str, bold: bool) -> float:
    f = _font(name, bool(bold))
    if f is None:
        return FALLBACK_EM.get(name, (0.65, 1.3))[1]
    asc, desc = f.getmetrics()
    return (asc + desc) / 100


def measure(tf, width_pt: float) -> tuple[float, int]:
    """(height in points, number of soft line wraps) for a filled text frame."""
    height, wraps = 0.0, 0
    for p in tf.paragraphs:
        pPr = p._p.pPr
        indent = int(pPr.get("marL", 0)) / 12700 if pPr is not None else 0
        avail = (width_pt - indent) * SAFETY
        lines, hard, line_w, line_h = 1, 1, 0.0, 0.0
        for child in p._p:
            if child.tag == qn("a:br"):
                lines, hard, line_w = lines + 1, hard + 1, 0.0
                continue
            if child.tag != qn("a:r"):
                continue
            r = _Run(child, p)
            name, size, bold = r.font.name or BODY_FONT, r.font.size.pt, bool(r.font.bold)
            line_h = max(line_h, size * _line_em(name, bold))
            for piece in re.findall(r"\S+\s*|\s+", r.text):
                word = _width(piece.rstrip(), name, bold, size)
                if line_w and line_w + word > avail:   # wrap before this word
                    lines, line_w = lines + 1, 0.0
                    piece = piece.lstrip()
                    word = _width(piece.rstrip(), name, bold, size)
                if word > avail:   # a word wider than the box breaks mid-word
                    extra = int(word // avail)
                    lines, word = lines + extra, word - extra * avail
                line_w += word + _width(piece[len(piece.rstrip()):], name, bold, size)
        before = p.space_before.pt if p.space_before is not None else 0
        after = p.space_after.pt if p.space_after is not None else 0
        height += lines * line_h + before + after
        wraps += lines - hard
    return height, wraps


def fits(shape, wrap: bool = True) -> bool:
    """Does the text, at its current sizes, fit the shape (without wrapping, if wrap=False)?"""
    tf = shape.text_frame
    w = (shape.width - tf.margin_left - tf.margin_right) / 12700
    h = (shape.height - tf.margin_top - tf.margin_bottom) / 12700
    height, wraps = measure(tf, w)
    return height <= h * SAFETY and (wrap or wraps == 0)


def fitted(candidates, make):
    """Add a text box with make(*args) for each candidate until one fits.

    candidates: [(args, wrap_allowed), ...], best first; the last one always stays."""
    for n, (args, wrap) in enumerate(candidates):
        shape = make(*args)
        if n == len(candidates) - 1 or fits(shape, wrap):
            return shape
        shape._element.getparent().remove(shape._element)


def ladder(sizes, wrap=True):
    return [((z,), wrap) for z in sizes]


def two_lines(text: str) -> str:
    """Break a long heading once: after a colon, else a comma, else the middle space."""
    for sep in (": ", ", "):
        if sep in text:
            cut = min((m.end() for m in re.finditer(re.escape(sep), text)),
                      key=lambda e: abs(e - len(text) / 2))
            return text[:cut].rstrip() + "\n" + text[cut:]
    spaces = [m.start() for m in re.finditer(" ", text)]
    if not spaces:
        return text
    cut = min(spaces, key=lambda e: abs(e - len(text) / 2))
    return text[:cut] + "\n" + text[cut + 1:]


# ---------- shapes ----------

def textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE   # shrinks the GTA's later edits
    tf.vertical_anchor = anchor
    return tb


def bullet(p, char="■"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.45))))
    pPr.set("indent", str(-Emu(Inches(0.35))))
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def rect(shapes, x, y, w, h, fill, line="ink", line_pt=3.0, shape=MSO_SHAPE.RECTANGLE):
    s = shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = C[fill]
    if line:
        s.line.color.rgb = C[line]
        s.line.width = Pt(line_pt)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


ADEC = "http://schemas.microsoft.com/office/drawing/2017/decorative"


def decorative(shape):
    """Mark a shape decorative, so accessibility checkers don't ask it for alt text."""
    cNvPr = shape._element.xpath("./*[1]/p:cNvPr")[0]
    ext = etree.SubElement(etree.SubElement(cNvPr, qn("a:extLst")), qn("a:ext"))
    ext.set("uri", "{C183D7F6-B498-43B3-948B-1728B52AA6E4}")
    etree.SubElement(ext, f"{{{ADEC}}}decorative", nsmap={"adec": ADEC}).set("val", "1")


def window(slide, x, y, w, h, label, bar, body="panel", label_color="ink"):
    """Grouped retro window frame; returns the body box (x, y, w, h)."""
    g = slide.shapes.add_group_shape()
    bar_h = Inches(0.55)
    rect(g.shapes, x + Inches(0.12), y + Inches(0.12), w, h, "ink", line=None)  # hard shadow
    rect(g.shapes, x, y, w, h, body)
    rect(g.shapes, x, y, w, bar_h, bar)
    for n, col in enumerate(("pink", "blue", "mint")):
        rect(g.shapes, x + Inches(0.25 + n * 0.33), y + Inches(0.16),
             Inches(0.22), Inches(0.22), col, line_pt=2.0)
    lt = g.shapes.add_textbox(x + Inches(1.35), y, w - Inches(1.5), bar_h).text_frame
    lt.vertical_anchor = MSO_ANCHOR.MIDDLE
    write(lt.paragraphs[0], label, 16, MONO_FONT, label_color, bold=True)
    g.name = f"Window {label}"
    decorative(g)
    pad = Inches(0.4)
    return x + pad, y + bar_h + Inches(0.25), w - 2 * pad, h - bar_h - Inches(0.5)


def sticker(slide, x, y, label: str, unit: str):
    """Pill-shaped unit badge like the web title slides; returns (group, height)."""
    size, pad, h = 18, Inches(0.3), Inches(0.5)
    w = Emu(int(_width(label, TITLE_FONT, False, size) * 12700)) + 2 * pad
    g = slide.shapes.add_group_shape()
    shadow = rect(g.shapes, x + Inches(0.06), y + Inches(0.06), w, h, "ink", line=None,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    pill = rect(g.shapes, x, y, w, h, UNIT[unit], line_pt=2.5,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    for s in (shadow, pill):
        s.adjustments[0] = 0.5
    tf = pill.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    write(tf.paragraphs[0], label, size, TITLE_FONT,
          "ink" if unit in DARK_TEXT_UNITS else "panel")
    g.name = f"Unit {label}"
    return g, h


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C[color]


def hidden_title(slide, text: str):
    """Fill the Title Only layout's title, then park it just off the slide.

    Screen readers, PowerPoint's Accessibility Checker and Canvas's Ally all
    read it; nobody sees it. (This is PowerPoint's own "hidden slide title".)"""
    for ph in list(slide.placeholders):
        if ph.placeholder_format.type not in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            ph._element.getparent().remove(ph._element)
    t = slide.shapes.title
    t.text_frame.text = plain(text)
    t.left, t.top, t.width, t.height = -W, 0, W, Inches(1)


def slide_title(s: dict) -> str:
    if "section" in s:
        return s["section"]
    if "quote" in s:
        return "Quote: " + plain(s.get("cite", ""))
    if "image" in s:
        if s.get("caption"):
            return s["caption"]
        alt = plain(s.get("alt", "")) or "Image"
        return alt if len(alt) <= 80 else alt[:80].rsplit(" ", 1)[0] + "…"
    if "points" in s:
        return s["title"]
    if "readings" in s:
        return s.get("title", "Readings")
    if "prompt" in s:
        return "Prompt: " + plain(s["title"]) if s.get("title") else "Prompt"
    if "question" in s:
        return "Discussion question"
    raise ValueError(f"unknown slide type: {s}")


def add_notes(slide, s):
    if s.get("notes"):
        slide.notes_slide.notes_text_frame.text = s["notes"].strip()


def fitted_picture(slide, src: Path, x, y, w, h, alt: str):
    im = Image.open(src)
    if max(im.size) > MAX_PX:
        im.thumbnail((MAX_PX, MAX_PX))
    buf = io.BytesIO()
    fmt = "PNG" if src.suffix.lower() == ".png" else "JPEG"
    (im if fmt == "PNG" else im.convert("RGB")).save(buf, fmt, quality=88)
    buf.seek(0)
    ratio = min(w / im.size[0], h / im.size[1])
    pw, ph = int(im.size[0] * ratio), int(im.size[1] * ratio)
    pic = slide.shapes.add_picture(buf, x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def bar_for(index: int) -> str:
    return f"bar{index % 4}"


def new_slide(prs, title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[5])   # Title Only
    hidden_title(slide, title)
    bg(slide, "paper")
    return slide


# ---------- slide types ----------

def render(prs, s: dict, index: int, unit: str):
    slide = new_slide(prs, slide_title(s))
    bar = bar_for(index + 1)
    if "section" in s:
        bg(slide, UNIT[unit])
        fg = "ink" if unit in DARK_TEXT_UNITS else "panel"
        text = two_lines(s["section"]) if ": " in s["section"] else s["section"]
        box_h = Inches(4.4)

        def make(size, sub_size):
            tb = textbox(slide, Inches(0.3), (H - box_h) // 2, W - Inches(0.6), box_h,
                         anchor=MSO_ANCHOR.MIDDLE)
            tf = tb.text_frame
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            write(tf.paragraphs[0], text, size, TITLE_FONT, fg)
            if s.get("sub"):
                p = tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                p.space_before = Pt(10)
                write(p, s["sub"], sub_size, MONO_FONT, fg)
            return tb
        fitted([((z, b), False) for z in (54, 48, 44, 40) for b in (28, 24, 22)]
               + [((44, 22), True)], make)
    elif "quote" in s:
        bx, by, bw, bh = window(slide, Inches(0.9), Inches(0.7), W - Inches(1.8),
                                H - Inches(1.4), s.get("label", "quote.txt"), bar)

        def make(size):
            tb = textbox(slide, bx, by, bw, bh - Inches(0.8), anchor=MSO_ANCHOR.MIDDLE)
            write(tb.text_frame.paragraphs[0], f"“{s['quote'].strip()}”", size)
            return tb
        fitted(ladder((28, 26, 24, 22, 20)), make)
        cite = f"[{s.get('cite', '')}]({s['url']})" if s.get("url") else s.get("cite", "")

        def make_cite(size):
            tb = textbox(slide, bx, by + bh - Inches(0.7), bw, Inches(0.6))
            tb.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
            write(tb.text_frame.paragraphs[0], f"— {cite}", size, MONO_FONT, "soft")
            return tb
        fitted(ladder((16, 14), wrap=False) + ladder((14,)), make_cite)
    elif "image" in s:
        src = REPO / s["image"]
        if not src.exists():
            raise FileNotFoundError(f"image not found: {s['image']}")
        cap_h = Inches(0.8) if (s.get("caption") or s.get("credit_url")) else 0
        bx, by, bw, bh = window(slide, Inches(0.9), Inches(0.4), W - Inches(1.8),
                                H - Inches(0.8) - cap_h, s.get("label", "image.png"), bar,
                                body="ink")
        fitted_picture(slide, src, bx, by, bw, bh, s.get("alt", ""))
        if cap_h:
            cap = s.get("caption", "")
            if s.get("credit_url"):
                cap = f"{cap} [source]({s['credit_url']})".strip()

            def make(size):
                tb = textbox(slide, Inches(0.9), H - cap_h - Inches(0.2), W - Inches(1.8),
                             cap_h)
                tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
                write(tb.text_frame.paragraphs[0], cap, size)
                return tb
            fitted(ladder((18, 16), wrap=False) + ladder((16, 14)), make)
    elif "prompt" in s:
        bx, by, bw, bh = window(slide, Inches(1.2), Inches(1.2), W - Inches(2.4),
                                H - Inches(2.4), "prompt.exe", bar, body="termbg")

        def make(size):
            tb = textbox(slide, bx, by, bw, bh, anchor=MSO_ANCHOR.MIDDLE)
            p = tb.text_frame.paragraphs[0]
            if s.get("title"):
                write(p, s["title"], 20, MONO_FONT, "bar3")
                p = tb.text_frame.add_paragraph()
                p.space_before = Pt(6)
            write(p, "> ", size, MONO_FONT, "yellow", bold=True)
            write(p, s["prompt"], size, MONO_FONT, "termfg")
            return tb
        fitted(ladder((28, 26, 24, 22, 20)), make)
    elif "question" in s:
        bx, by, bw, bh = window(slide, Inches(1.2), Inches(1.0), W - Inches(2.4),
                                H - Inches(2.0), s.get("label", "discuss.txt"), bar)

        def make(size):
            tb = textbox(slide, bx, by, bw, bh, anchor=MSO_ANCHOR.MIDDLE)
            write(tb.text_frame.paragraphs[0], s["question"], size, TITLE_FONT)
            if s.get("hint"):
                p = tb.text_frame.add_paragraph()
                p.space_before = Pt(12)
                write(p, s["hint"], 20, BODY_FONT, "soft")
            return tb
        fitted(ladder((32, 30, 28, 26, 24)), make)
    elif "points" in s or "readings" in s:
        title = s.get("title", "Readings") if "readings" in s else s["title"]

        def make_title(size):
            tb = textbox(slide, Inches(0.8), Inches(0.5), W - Inches(1.6), Inches(1.0))
            write(tb.text_frame.paragraphs[0], title, size, TITLE_FONT,
                  UNIT[unit] if unit not in DARK_TEXT_UNITS else "ink")
            return tb
        fitted(ladder((36, 32, 28), wrap=False) + ladder((28,)), make_title)
        items = s.get("points") or s["readings"]

        def make(scale):
            tb = textbox(slide, Inches(0.8), Inches(1.7), W - Inches(1.6), H - Inches(2.3))
            for n, item in enumerate(items):
                p = tb.text_frame.paragraphs[0] if n == 0 else tb.text_frame.add_paragraph()
                bullet(p)
                p.space_after = Pt(12)
                if isinstance(item, dict):
                    write(p, item["cite"], round(22 * scale))
                    if item.get("what"):
                        write(p, f" — {item['what']}", round(18 * scale), BODY_FONT, "soft")
                else:
                    write(p, item, round(24 * scale))
            return tb
        fitted(ladder((1.0, 0.92, 0.84, 0.75)), make)
    else:
        raise ValueError(f"unknown slide type: {s}")
    add_notes(slide, s)


def title_slide(prs, meta: dict):
    c = meta["canvas"]
    slide = new_slide(prs, c["module"])
    label, rest = (t.strip() for t in c["module"].split(":", 1))
    start, due = c["week_start"], c["due"]
    dates = f"{start:%b} {start.day} – {due:%b} {due.day}, {due.year}".upper()
    bx, by, bw, bh = window(slide, Inches(1.0), Inches(1.0), W - Inches(2.0),
                            H - Inches(2.0), "ENG3861.EXE", "pink", label_color="panel")
    badge, badge_h = sticker(slide, bx + Inches(0.1), by,
                             UNIT_LABELS[c["unit"]].upper(), c["unit"])
    gap = Inches(0.2)

    def make(name, size):
        tb = textbox(slide, bx, by + badge_h + gap, bw, bh - badge_h - gap)
        tf = tb.text_frame
        write(tf.paragraphs[0], label, 24, MONO_FONT, "soft")
        write(tf.add_paragraph(), name, size, TITLE_FONT)
        write(tf.add_paragraph(), f"> ENG 3861 // {dates}_", 20, MONO_FONT, "soft")
        return tb
    tb = fitted([((rest, z), False) for z in (48, 44, 40)]
                + [((two_lines(rest), z), False) for z in (48, 44, 40, 36)]
                + [((rest, 36), True)], make)
    # centre the sticker and the text below it as one block in the window body
    tf = tb.text_frame
    text_h = Emu(int(measure(tf, (bw - tf.margin_left - tf.margin_right) / 12700)[0] * 12700))
    block = badge_h + gap + text_h + tf.margin_top + tf.margin_bottom
    shift = max(0, (bh - block) // 2)
    for sh in badge.shapes:
        sh.top += shift
    badge._element.recalculate_extents()
    tb.top += shift
    tb.height = by + bh - tb.top


def build_deck(stem: str, slides: list[dict], meta: dict) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    title_slide(prs, meta)
    for i, s in enumerate(slides, start=1):
        render(prs, s, i, meta["canvas"]["unit"])
    cp = prs.core_properties
    cp.title = meta["canvas"]["module"]
    cp.author = "ENG 3861"
    cp.last_modified_by = ""
    cp.revision = 1
    cp.created = cp.modified = FIXED_DATE
    return prs


ZIP_DATE = (2027, 1, 1, 0, 0, 0)


def save_deck(prs: Presentation, out: Path) -> None:
    """Save with every zip entry stamped ZIP_DATE, not the build time, so an
    unchanged deck rebuilds byte for byte (same sha, no re-upload)."""
    buf = io.BytesIO()
    prs.save(buf)
    fixed = io.BytesIO()
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(fixed, "w") as dst:
        for info in src.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=ZIP_DATE)
            entry.compress_type = info.compress_type
            entry.external_attr = info.external_attr
            dst.writestr(entry, src.read(info.filename))
    Path(out).write_bytes(fixed.getvalue())


def main() -> int:
    decks = yaml.safe_load((HERE / "decks.yml").read_text(encoding="utf-8"))
    for stem, slides in decks.items():
        meta = frontmatter.load(str(COURSE / f"{stem}.md")).metadata
        out = HERE / f"{stem}.pptx"
        save_deck(build_deck(stem, slides, meta), out)
        print(f"{stem}.pptx: {len(slides) + 1} slides, {out.stat().st_size // 1024} KB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
